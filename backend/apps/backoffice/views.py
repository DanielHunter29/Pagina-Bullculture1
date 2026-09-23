"""Vistas del backoffice (dashboard y reportes). Solo personal (staff)."""
from io import BytesIO

from django.contrib.admin.views.decorators import staff_member_required
from django.http import HttpResponse
from django.shortcuts import render

from apps.common.money import format_cop
from apps.orders.models import Order

from .reports import (
    expenses_summary,
    parse_range,
    product_profitability,
    sales_summary,
)


@staff_member_required
def dashboard(request):
    from .reports import dashboard_metrics

    metrics = dashboard_metrics()
    context = {
        "metrics": metrics,
        "recent_orders": Order.objects.order_by("-created_at")[:10],
        "title": "Dashboard · BULLCULTURE",
    }
    return render(request, "backoffice/dashboard.html", context)


def _report_context(request):
    d_from, d_to = parse_range(request.GET.get("desde"), request.GET.get("hasta"))
    return (
        d_from,
        d_to,
        sales_summary(d_from, d_to),
        product_profitability(d_from, d_to),
        expenses_summary(d_from, d_to),
    )


@staff_member_required
def export_sales_xlsx(request):
    import openpyxl

    d_from, d_to, sales, profit, expenses = _report_context(request)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Resumen"
    net = sales["revenue"] - expenses["total"]
    for row in [
        ("Reporte BULLCULTURE", ""),
        ("Desde", d_from.isoformat()),
        ("Hasta", d_to.isoformat()),
        ("Pedidos pagados", sales["orders_count"]),
        ("Unidades vendidas", sales["units_sold"]),
        ("Ingresos", float(sales["revenue"])),
        ("Gastos", float(expenses["expenses"])),
        ("Facturas proveedor", float(expenses["supplier_invoices"])),
        ("Utilidad neta", float(net)),
    ]:
        ws.append(row)

    wp = wb.create_sheet("Rentabilidad")
    wp.append(["Producto", "Unidades", "Ingresos", "Costo", "Utilidad"])
    for r in profit:
        wp.append([
            r["product_name"],
            r["units"],
            float(r["revenue"]),
            float(r["cost"]),
            float(r["profit"]),
        ])

    wv = wb.create_sheet("Ventas")
    wv.append(["Referencia", "Cliente", "Ciudad", "Total", "Estado pago", "Estado envío", "Pagado"])
    for o in Order.objects.filter(
        payment_status=Order.PaymentStatus.APPROVED, paid_at__date__range=(d_from, d_to)
    ).order_by("-paid_at"):
        wv.append([
            o.reference,
            o.customer_name,
            o.shipping_city,
            float(o.total),
            o.get_payment_status_display(),
            o.get_fulfillment_status_display(),
            o.paid_at.strftime("%Y-%m-%d %H:%M") if o.paid_at else "",
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    resp = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = (
        f'attachment; filename="reporte_{d_from}_{d_to}.xlsx"'
    )
    return resp


@staff_member_required
def export_sales_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    d_from, d_to, sales, profit, expenses = _report_context(request)
    net = sales["revenue"] - expenses["total"]

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title="Reporte BULLCULTURE")
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("BULLCULTURE — Reporte de ventas", styles["Title"]),
        Paragraph(f"Del {d_from} al {d_to}", styles["Normal"]),
        Spacer(1, 0.5 * cm),
    ]

    money = format_cop

    summary = [
        ["Pedidos pagados", str(sales["orders_count"])],
        ["Unidades vendidas", str(sales["units_sold"])],
        ["Ingresos", money(sales["revenue"])],
        ["Gastos + facturas", money(expenses["total"])],
        ["Utilidad neta", money(net)],
    ]
    t = Table(summary, colWidths=[8 * cm, 6 * cm])
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#0A0E14")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
    ]))
    elements += [t, Spacer(1, 0.6 * cm), Paragraph("Rentabilidad por producto", styles["Heading2"])]

    data = [["Producto", "Uds.", "Ingresos", "Costo", "Utilidad"]]
    for r in profit:
        data.append([
            r["product_name"],
            str(r["units"]),
            money(r["revenue"]),
            money(r["cost"]),
            money(r["profit"]),
        ])
    if len(data) == 1:
        data.append(["Sin ventas en el rango", "", "", "", ""])
    pt = Table(data, colWidths=[6 * cm, 2 * cm, 3 * cm, 3 * cm, 3 * cm])
    pt.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6C93B6")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    elements.append(pt)

    doc.build(elements)
    buffer.seek(0)
    resp = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="reporte_{d_from}_{d_to}.pdf"'
    return resp
